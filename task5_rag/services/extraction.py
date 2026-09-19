"""
Extracts raw text from uploaded documents, page by page where possible.
Returns a list of (page_number, text) tuples so downstream chunking can
retain page-level metadata.
"""
import re
from pathlib import Path
from typing import List, Tuple

import pdfplumber
from docx import Document as DocxDocument


def clean_text(text: str) -> str:
    """Collapse excess whitespace/newlines left over from PDF/DOCX extraction."""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(path: Path) -> List[Tuple[int, str]]:
    pages = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            raw = page.extract_text() or ""
            pages.append((i, clean_text(raw)))
    return pages


def extract_docx(path: Path) -> List[Tuple[int, str]]:
    """DOCX has no native page concept, so the whole document is treated as page 1."""
    doc = DocxDocument(path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)
    full_text = clean_text("\n".join(parts))
    return [(1, full_text)]


def extract_txt(path: Path) -> List[Tuple[int, str]]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    return [(1, clean_text(raw))]


def extract_text(path: Path, file_type: str) -> List[Tuple[int, str]]:
    """
    file_type: one of 'pdf', 'docx', 'txt'
    Returns list of (page_number, page_text). Empty pages are kept out.
    """
    if file_type == "pdf":
        pages = extract_pdf(path)
    elif file_type == "docx":
        pages = extract_docx(path)
    elif file_type == "txt":
        pages = extract_txt(path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")

    return [(p, t) for p, t in pages if t.strip()]
